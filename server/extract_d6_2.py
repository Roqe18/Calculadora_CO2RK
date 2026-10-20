import docx
import os

doc_path = "D6.2_CuantificacionBiomasaVivaIDAF.docx"
doc = docx.Document(doc_path)

output_file = "D6.2_contenido_extraido.txt"

with open(output_file, "w", encoding="utf-8") as f:
    f.write(f"=== DOCUMENTO: {doc_path} ===\n\n")
    
    # 1. Parrafos y Encabezados
    f.write("--- PARRAFOS Y SECCIONES ---\n")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            style = p.style.name if p.style else "Normal"
            f.write(f"[{style}] {text}\n")
            
    # 2. Tablas
    f.write("\n\n--- TABLAS DEL DOCUMENTO ---\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n[TABLA {t_idx + 1}] ({len(table.rows)} filas x {len(table.columns)} columnas)\n")
        for r_idx, row in enumerate(table.rows):
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            f.write(f"  Fila {r_idx + 1}: {' | '.join(row_cells)}\n")

print(f"Contenido extraido con exito en {output_file}")
